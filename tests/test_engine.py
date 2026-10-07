"""Regression tests for the FP&A engine (standard library unittest; run: python -m unittest discover -s tests).

Every test runs in its own temporary copy of config/ and data/, so tests never touch your real files.
Tests are grouped by workflow stage and by the original bugs they guard against.
"""
import json
import os
import shutil
import sys
import tempfile
import unittest
import warnings
from pathlib import Path
from unittest import mock

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook

os.environ["FPA_QUIET"] = "1"                      # silence stage logging during tests
warnings.filterwarnings("ignore", category=DeprecationWarning)
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine import formatting as fmt  # noqa: E402
from engine import llm_narrative  # noqa: E402
from engine.anonymizer import Redactor  # noqa: E402
from engine.config import CARD_MAX_CHARS, SKILL_MAX_CHARS, Paths  # noqa: E402
from engine.deck import build_deck, check_deck  # noqa: E402
from engine.history import HistoryError, lock_period, restate_period  # noqa: E402
from engine.ingestion import load_column_map, load_trial_balance, normalise_period, read_table  # noqa: E402
from engine.kpi import kpi_status, score_kpis  # noqa: E402
from engine.mapping import apply_statutory_mapping  # noqa: E402
from engine.pipeline import approve_all, lock_after_release, run_pipeline, signoff_action  # noqa: E402
from engine.revision_log import build_revision_log  # noqa: E402
from engine.signoff import (FINANCE_REVIEW, FPA_EDIT, RELEASED, SignoffError, approve, new_state, pending_roles,  # noqa: E402
                            record_edit, reject, resubmit)
from engine.llm_narrative import check_ollama_connection, extract_json_object, model_is_available, parse_ai_json  # noqa: E402
from engine.validation import check_forbidden_terms, check_numbers, check_slide_spec  # noqa: E402
from engine.variance_ratios import (compute_comparisons, compute_ratios, days_in, favorability, flag_material,  # noqa: E402
                                    period_add, quarter_to_date, year_to_date)

SEEDED = ["Orchid Bay", "Sunrise Logistics", "CloudNimbus", "Meridian Health", "Jane Tan", "jane.tan@", "ap@cloudnimbus", "088-123456-7", "+65 8123"]
BU_NAMES = ["Platform", "Enterprise", "SMB"]


def make_workspace() -> Path:
    """Copy config/ and data/ into a fresh temporary project folder."""
    ws = Path(tempfile.mkdtemp())
    for folder in ("config", "data"):
        shutil.copytree(ROOT / folder, ws / folder)
    (ws / "outputs").mkdir()
    return ws


def edit_erp(path: Path, change) -> None:
    """Load an ERP export (all text), apply `change(df)` in place and save it back."""
    df = pd.read_excel(path, dtype=str)
    change(df)
    df.to_excel(path, index=False)


def set_amount(df, code, bu, delta):
    mask = (df["Account Code"] == code) & (df["Business Unit"] == bu)
    df.loc[mask, "Balance"] = str(float(df.loc[mask, "Balance"].iloc[0]) + delta)


class WorkspaceCase(unittest.TestCase):
    """Throw-away workspace plus quick pipeline runners for the Day-3 and Day-4 exports."""

    def setUp(self):
        self.ws = make_workspace()
        self.paths = Paths(self.ws)
        self.day3, self.day4 = self.paths.samples / "day3_trial_balance.xlsx", self.paths.samples / "day4_trial_balance.xlsx"

    def tearDown(self):
        shutil.rmtree(self.ws, ignore_errors=True)

    def run_day3(self, **kw):
        kw.setdefault("write_outputs", False); kw.setdefault("build_decks", False)
        return run_pipeline(tb_file=self.day3, use_mock=kw.pop("use_mock", True), base_dir=self.ws, **kw)

    def run_day4(self, **kw):
        kw.setdefault("write_outputs", False); kw.setdefault("build_decks", False)
        return run_pipeline(tb_file=self.day4, previous_tb_file=kw.pop("previous", self.day3), use_mock=kw.pop("use_mock", True), base_dir=self.ws, **kw)


# ======================================================================= helpers and formulas
class TestPeriodHelpers(unittest.TestCase):
    def test_period_arithmetic(self):
        self.assertEqual(period_add("2026-01", -1), "2025-12")
        self.assertEqual(period_add("2026-06", -12), "2025-06")
        self.assertEqual(quarter_to_date("2026-06"), ["2026-04", "2026-05", "2026-06"])
        self.assertEqual(len(year_to_date("2026-06")), 6)

    def test_days_use_real_calendar(self):
        self.assertEqual(days_in(["2026-06"]), 30)
        self.assertEqual(days_in(quarter_to_date("2026-06")), 91)
        self.assertEqual(days_in(["2024-02"]), 29)


class TestVarianceLogic(unittest.TestCase):
    def test_favorability_depends_on_category(self):
        self.assertEqual(favorability("Revenue", 35000), "Favorable")
        self.assertEqual(favorability("COGS", 35000), "Unfavorable")
        self.assertEqual(favorability("OpEx", -5000), "Favorable")
        self.assertEqual(favorability("Revenue", 0), "Neutral")

    def test_materiality_is_two_sided(self):
        df = pd.DataFrame({"Var_Amt": [30000, 10000, 30000, 30000], "Var_Pct": [0.02, 0.50, 0.10, np.nan]})
        self.assertEqual(flag_material(df, 25000, 0.05)["Material"].tolist(), [False, False, True, True])

    def test_zero_comparator_gives_nan_not_absurd_percentage(self):
        """Original bug: replace(0, 1) produced huge fake percentages when the budget was zero."""
        rows = [("Actual", "2026-06", "All", "Product Revenue", "Revenue", 100000.0), ("Actual", "2026-06", "All", "New Line", "Revenue", 30000.0),
                ("Budget", "2026-06", "All", "Product Revenue", "Revenue", 90000.0)]
        agg = pd.DataFrame(rows, columns=["Scenario", "Period", "BU", "Line", "Category", "Amount"])
        cmp_df, _ = compute_comparisons(agg, "2026-06", "statutory", 25000, 0.05)
        row = cmp_df[(cmp_df["Horizon"] == "BvA_Month") & (cmp_df["Line"] == "New Line")].iloc[0]
        self.assertTrue(np.isnan(row["Var_Pct"]) and row["Material"])

    def test_zero_denominators_in_ratios_do_not_crash(self):
        agg = pd.DataFrame([("Actual", "2026-06", "Corporate", "Accounts Receivable", "Asset", 0.0)], columns=["Scenario", "Period", "BU", "Line", "Category", "Amount"])
        r = compute_ratios(agg, "2026-06")
        self.assertEqual((r["DSO_Days"], r["DPO_Days"], r["Gross_Margin_%"]), (0.0, 0.0, 0.0))

    def test_quarter_ratio_uses_91_days(self):
        rows = [("Actual", p, "All", "Product Revenue", "Revenue", 100000.0) for p in quarter_to_date("2026-06")]
        rows += [("Actual", p, "All", "Cost of Revenue", "COGS", 40000.0) for p in quarter_to_date("2026-06")]
        rows += [("Actual", "2026-06", "Corporate", "Accounts Receivable", "Asset", 300000.0), ("Actual", "2026-06", "Corporate", "Accounts Payable", "Liability", 80000.0)]
        r = compute_ratios(pd.DataFrame(rows, columns=["Scenario", "Period", "BU", "Line", "Category", "Amount"]), "2026-06", "QTD")
        self.assertEqual(r["Days_In_Period"], 91)
        self.assertAlmostEqual(r["DSO_Days"], 91.0, places=2)


# ======================================================================= stage 1-3
class TestIngestionAndMapping(WorkspaceCase):
    def test_clean_run_succeeds_and_all_stages_pass(self):
        r = self.run_day4()
        self.assertEqual(r["status"], "SUCCESS", r["error"])
        self.assertEqual([s["status"] for s in r["stages"]], ["PASS"] * 8 + ["SKIPPED", "PASS"])
        self.assertEqual(r["warnings"], [])

    def test_erp_export_needs_no_period_or_version_columns(self):
        cols = list(pd.read_excel(self.day3).columns)
        self.assertEqual(cols, ["Account Code", "Account Name", "Business Unit", "Balance", "Journal Description"])
        self.assertEqual(self.run_day3()["status"], "SUCCESS")

    def test_thousands_separators_in_amounts_are_accepted(self):
        edit_erp(self.day3, lambda df: df.__setitem__("Balance", df["Balance"].map(lambda v: f"{float(v):,.0f}")))
        self.assertEqual(self.run_day3()["status"], "SUCCESS")

    def test_stale_file_with_wrong_period_column_is_rejected(self):
        edit_erp(self.day3, lambda df: df.__setitem__("Period", "2026-05"))
        r = self.run_day3()
        self.assertEqual(r["status"], "FAILED"); self.assertIn("stale", r["error"])

    def test_export_identical_to_last_month_is_rejected(self):
        hist = pd.read_csv(self.paths.history, dtype={"ERP_Account_Code": str})
        may = hist[hist["Period"] == "2026-05"]
        pd.DataFrame({"Account Code": may["ERP_Account_Code"], "Account Name": "x", "Business Unit": may["BU"], "Balance": may["Amount"]}).to_excel(self.day3, index=False)
        r = self.run_day3()
        self.assertEqual(r["status"], "FAILED"); self.assertIn("last month", r["error"])

    def test_unmapped_account_blocks_the_run_instead_of_dropping_silently(self):
        """Original bug: unmapped accounts disappeared from totals while only a warning printed."""
        def add_row(df):
            df.loc[len(df)] = ["9999-99", "New GL", "Platform", "5000", "new account"]
        edit_erp(self.day4, add_row)
        r = self.run_day4()
        self.assertEqual(r["status"], "BLOCKED"); self.assertIn("unmapped", r["error"])
        self.assertTrue(any("9999-99" in w for w in r["warnings"]))

    def test_effective_dated_mapping(self):
        mapping = pd.DataFrame({"ERP_Account_Code": ["4000-01"] * 2, "Std_Category": ["Revenue"] * 2, "Std_Line": ["Old", "New"], "Effective_From": ["2025-01", "2026-01"]})
        fact = pd.DataFrame({"Scenario": "Actual", "Period": ["2024-12", "2025-12", "2026-01"], "ERP_Account_Code": "4000-01", "BU": "P", "Amount": 1.0})
        self.assertEqual(apply_statutory_mapping(fact, mapping).sort_values("Period")["Line"].tolist(), ["UNMAPPED", "Old", "New"])

    def test_invalid_reclass_percent_fails_clearly(self):
        path = self.ws / "config" / "mgmt_reclass_rules.csv"
        path.write_text(path.read_text().replace(",0.20,", ",1.50,"))
        r = self.run_day4()
        self.assertEqual(r["status"], "FAILED"); self.assertIn("between 0 and 1", r["error"])

    def test_malformed_config_file_gives_clear_message(self):
        path = self.ws / "config" / "exclusion_rules.csv"
        path.write_text(path.read_text() + "Investor,line,X,a, b, c, d\n")
        r = self.run_day4()
        self.assertEqual(r["status"], "FAILED"); self.assertIn("exclusion_rules.csv", r["error"])

    def test_missing_history_month_skips_horizon_instead_of_zero_filling(self):
        hist = pd.read_csv(self.paths.history, dtype=str)
        hist[hist["Period"] != "2026-03"].to_csv(self.paths.history, index=False)
        r = self.run_day4()
        self.assertEqual(r["status"], "SUCCESS")
        self.assertTrue(any("2026-03" in w for w in r["warnings"]) and any("QoQ" in w and "skipped" in w for w in r["warnings"]))
        pl = r["_frames"]["stat_cmp"]
        self.assertNotIn("QoQ", set(pl[pl["Statement"] == "P&L"]["Horizon"]))

    def test_date_converted_period_is_repaired_with_warning(self):
        w = []
        self.assertEqual(normalise_period("2026-06-01 00:00:00", w), "2026-06")
        self.assertEqual(normalise_period(pd.Timestamp("2026-06-15"), w), "2026-06")
        self.assertEqual(len(w), 2)
        with self.assertRaises(Exception):
            normalise_period("June", [])


class TestCalculations(WorkspaceCase):
    def setUp(self):
        super().setUp()
        self.r = self.run_day4()
        self.cmp = self.r["_frames"]["stat_cmp"]

    def row(self, horizon, line, statement="P&L", df=None):
        df = self.cmp if df is None else df
        return df[(df["Statement"] == statement) & (df["Horizon"] == horizon) & (df["Line"] == line) & (df["BU"] == "All")].iloc[0]

    def test_reclass_moves_gross_margin_but_not_operating_profit(self):
        bridge = {b["Step"]: b["Amount"] for b in self.r["bridge"]["table"]}
        self.assertEqual(bridge["Difference (must be zero)"], 0)
        self.assertAlmostEqual(bridge["Less reclass R1: 20% of engineering payroll apportioned to cost of revenue"], -32000)
        self.assertAlmostEqual(bridge["Less reclass R2: 30% of admin overhead apportioned to cost of revenue"], -18900)
        self.assertLess(self.r["management_ratios"]["Adjusted_Gross_Margin_%"], self.r["ratios"]["Gross_Margin_%"])

    def test_variances_match_independent_calculation(self):
        hist = pd.read_csv(self.paths.history, dtype={"Period": str, "ERP_Account_Code": str})
        rev = lambda periods: hist[hist["Period"].isin(periods) & hist["ERP_Account_Code"].isin(["4000-01", "4000-02"])]["Amount"].sum()  # noqa: E731
        self.assertAlmostEqual(self.row("MoM", "Total Revenue")["Comparator"], rev(["2026-05"]))
        self.assertAlmostEqual(self.row("YoY", "Total Revenue")["Comparator"], rev(["2025-06"]))
        self.assertAlmostEqual(self.row("QoQ", "Total Revenue")["Comparator"], rev(["2026-01", "2026-02", "2026-03"]))
        self.assertAlmostEqual(self.row("QoQ", "Total Revenue")["Actual"], rev(["2026-04", "2026-05"]) + 608000)
        cost = self.row("BvA_Month", "Cost of Revenue")
        self.assertAlmostEqual(cost["Var_Amt"], 49500); self.assertEqual(cost["Favorability"], "Unfavorable")

    def test_working_capital_formulas(self):
        r = self.r["ratios"]
        self.assertAlmostEqual(r["DSO_Days"], round(908000 / 608000 * 30, 2))
        self.assertAlmostEqual(r["DPO_Days"], round(172000 / 199500 * 30, 2))
        self.assertAlmostEqual(r["Working_Capital_Drag_Days"], r["DSO_Days"] - r["DPO_Days"], places=2)

    def test_materiality_counts_follow_thresholds(self):
        self.assertEqual(self.run_day4(abs_threshold=1_000_000)["material_variances_count"], 0)
        self.assertGreater(self.run_day4(abs_threshold=1000, rel_threshold=0.0)["material_variances_count"], self.r["material_variances_count"])


# ======================================================================= balance sheet and cash flow
class TestStatements(WorkspaceCase):
    def test_balance_sheet_balances_and_cash_flow_reconciles_for_every_scenario_month(self):
        from engine.statements import balance_check, cash_reconciliation
        f = self.run_day4()["_frames"]
        self.assertLess(balance_check(f["bs"])["Difference"].abs().max(), 0.01)
        rec = cash_reconciliation(f["bs"], f["cf"])
        self.assertLess(rec["Difference"].abs().max(), 0.01)
        self.assertEqual(set(rec["Scenario"]), {"Actual", "Budget", "Forecast"})

    def test_cash_flow_matches_change_in_cash(self):
        f = self.run_day4()["_frames"]
        cash = lambda p: f["bs"][(f["bs"]["Scenario"] == "Actual") & (f["bs"]["Period"] == p) & (f["bs"]["Line"] == "Cash")]["Amount"].iloc[0]  # noqa: E731
        net = f["cf"][(f["cf"]["Scenario"] == "Actual") & (f["cf"]["Period"] == "2026-06") & (f["cf"]["Line"] == "Net Cash Flow")]["Amount"].iloc[0]
        self.assertAlmostEqual(net, cash("2026-06") - cash("2026-05"))

    def test_statement_comparisons_exist_with_balance_sheet_point_in_time_horizons(self):
        cmp_df = self.run_day4()["_frames"]["stat_cmp"]
        self.assertEqual(set(cmp_df[cmp_df["Statement"] == "BS"]["Horizon"]), {"BvA_Month", "RFvA_Month", "MoM", "QoQ", "YoY"})
        self.assertEqual(len(set(cmp_df[cmp_df["Statement"] == "CF"]["Horizon"])), 9)
        cash = cmp_df[(cmp_df["Statement"] == "BS") & (cmp_df["Horizon"] == "BvA_Month") & (cmp_df["Line"] == "Cash")].iloc[0]
        self.assertEqual(cash["Favorability"], "Unfavorable")                # cash below budget

    def test_unbalanced_balance_sheet_blocks_the_run(self):
        edit_erp(self.day4, lambda df: set_amount(df, "1100-00", "Corporate", 5000))     # receivable up with no offsetting entry
        r = self.run_day4()
        self.assertEqual(r["status"], "BLOCKED"); self.assertIn("balance sheet difference", r["error"])

    def test_broken_retained_earnings_roll_blocks_on_cash_reconciliation(self):
        edit_erp(self.day4, lambda df: (set_amount(df, "3100-00", "Corporate", 7000), set_amount(df, "1000-00", "Corporate", 7000)))
        r = self.run_day4()
        self.assertEqual(r["status"], "BLOCKED"); self.assertIn("cash reconciliation difference", r["error"])


class TestKpis(WorkspaceCase):
    def test_status_bands(self):
        self.assertEqual(kpi_status(72, 70, "Higher", 5), "On target")
        self.assertEqual(kpi_status(67.2, 70, "Higher", 5), "Watch")          # inside the 3.5-point band
        self.assertEqual(kpi_status(60, 70, "Higher", 5), "Off target")
        self.assertEqual(kpi_status(44, 42, "Lower", 5), "Watch")
        self.assertEqual(kpi_status(None, 18, "Higher", 10), "On target")      # runway None = cash generative

    def test_unknown_kpi_target_is_ignored_with_warning(self):
        w = []
        df = score_kpis({"Gross_Margin_%": 70.0}, pd.DataFrame([{"KPI": "Nonsense", "Direction": "Higher", "Target": "1", "Tolerance_Pct": "5", "Visible_To": "Management"}]), w)
        self.assertTrue(df.empty and any("Nonsense" in x for x in w))

    def test_scorecard_values_and_audience_visibility(self):
        r = self.run_day4(write_outputs=False)
        names = {k["KPI"]: k for k in r["kpis"]}
        self.assertAlmostEqual(names["Gross_Margin_%"]["Value"], r["ratios"]["Gross_Margin_%"])
        self.assertEqual(names["Cash_Balance"]["Value"], 2153700.0)
        packs = r["_frames"]["packs"]
        self.assertNotIn("Cash balance", packs["Investor"]["kpis"]["Label"].tolist())      # investors chose margins, growth, working capital
        self.assertIn("Cash balance", packs["Management"]["kpis"]["Label"].tolist())

    def test_adjusted_gross_margin_is_a_distinct_kpi_visible_to_management_only_and_lower_than_statutory(self):
        """Bug: the dashboard's KPI scorecard looked identical across audiences because Adjusted Gross Margin
        was never a KPI at all (only a row in the separate 'ratios' table). It is now a KPI in its own right."""
        r = self.run_day4(write_outputs=False)
        names = {k["KPI"]: k for k in r["kpis"]}
        self.assertIn("Adjusted_Gross_Margin_%", names)
        self.assertLess(names["Adjusted_Gross_Margin_%"]["Value"], names["Gross_Margin_%"]["Value"])
        packs = r["_frames"]["packs"]
        mgmt_labels = packs["Management"]["kpis"]["Label"].tolist()
        self.assertIn("Adjusted gross margin (management basis)", mgmt_labels)
        self.assertIn("Gross margin (statutory)", mgmt_labels)                              # both shown side by side
        for audience in ("Board", "Investor"):
            self.assertNotIn("Adjusted gross margin (management basis)", packs[audience]["kpis"]["Label"].tolist())

    def test_adjusted_opex_pct_is_a_distinct_kpi_lower_than_statutory_because_reclass_moves_cost_out_of_opex(self):
        """Bug: Opex % of revenue was computed from the statutory basis only, so it looked identical for every
        audience even though the management reclass moves part of payroll/admin out of OpEx into COGS, which
        should make the management-basis figure lower."""
        r = self.run_day4(write_outputs=False)
        names = {k["KPI"]: k for k in r["kpis"]}
        self.assertIn("Adjusted_Opex_%_of_Revenue", names)
        self.assertLess(names["Adjusted_Opex_%_of_Revenue"]["Value"], names["Opex_%_of_Revenue"]["Value"])
        packs = r["_frames"]["packs"]
        self.assertIn("Adjusted opex % of revenue (management basis)", packs["Management"]["kpis"]["Label"].tolist())
        for audience in ("Board", "Investor"):
            self.assertNotIn("Adjusted opex % of revenue (management basis)", packs[audience]["kpis"]["Label"].tolist())

    def test_runway_text_is_not_truncated_when_cash_generative(self):
        r = self.run_day4(write_outputs=False)
        runway = [k for k in r["kpis"] if k["KPI"] == "Runway_Months"][0]
        self.assertEqual(runway["Value_Text"], "Cash generative")

    def test_runway_is_none_when_cash_generative(self):
        self.assertIsNone([k for k in self.run_day4()["kpis"] if k["KPI"] == "Runway_Months"][0]["Value"])
        self.assertTrue(all(k["Value"] is not None for k in self.run_day4()["kpis"] if k["KPI"] != "Runway_Months"))


# ======================================================================= stage 4: revision log
class TestRevisionLog(WorkspaceCase):
    def test_day4_changes_detected_and_explained_with_minimal_effort(self):
        log = self.run_day4()["_frames"]["revision_log"]
        self.assertEqual(len(log), 6)
        self.assertEqual(log["Status"].value_counts().to_dict(), {"Explained (ledger memo)": 5, "Explained (manual)": 1})
        self.assertAlmostEqual(log["Change"].sum(), -22000 - 22000 + 14500 + 14500)

    def test_material_change_without_any_explanation_blocks(self):
        def blank_memo(df):
            df.loc[(df["Account Code"] == "5000-10") & (df["Business Unit"] == "Platform"), "Journal Description"] = None
        edit_erp(self.day4, blank_memo)
        r = self.run_day4()
        self.assertEqual(r["status"], "BLOCKED"); self.assertIn("5000-10/Platform", r["error"])

    def test_small_change_without_explanation_is_auto_logged(self):
        edit_erp(self.day4, lambda df: (set_amount(df, "6000-90", "SMB", 500), set_amount(df, "2200-00", "Corporate", 500)))   # balanced entry
        r = self.run_day4()
        self.assertEqual(r["status"], "SUCCESS", r["error"])
        self.assertIn("Auto-logged (below threshold)", set(r["_frames"]["revision_log"]["Status"]))

    def test_unbalanced_small_change_is_still_caught_by_the_balance_sheet_control(self):
        edit_erp(self.day4, lambda df: set_amount(df, "6000-90", "SMB", 500))
        self.assertEqual(self.run_day4()["status"], "BLOCKED")

    def test_manual_reason_overrides_memo(self):
        row = self.run_day4()["_frames"]["revision_log"].query("ERP_Account_Code == '4000-01'").iloc[0]
        self.assertEqual((row["Status"], row["Adjustment_Ref"]), ("Explained (manual)", "ADJ-01"))

    def test_line_added_or_removed_between_runs_is_caught(self):
        cols = ["Period", "ERP_Account_Code", "BU", "ERP_Description", "Amount", "Memo"]
        a = pd.DataFrame([["2026-06", "4000-01", "P", "Rev", 100.0, ""]], columns=cols)
        b = pd.DataFrame([["2026-06", "5000-10", "P", "Cost", 50.0, ""]], columns=cols)
        log = build_revision_log(a, b, pd.DataFrame(columns=["Period", "ERP_Account_Code", "BU", "Adjustment_Ref", "Reason", "Prepared_By"]), 10)
        self.assertEqual(sorted(log["Change"]), [-100.0, 50.0]); self.assertEqual(set(log["Status"]), {"UNEXPLAINED"})

    def test_run_one_skips_revision_log(self):
        r = self.run_day3()
        self.assertEqual(r["stages"][3]["status"], "SKIPPED")
        self.assertIn("DRAFT run 1", r["narratives"]["Management"])


class TestRunNumbering(WorkspaceCase):
    def test_same_file_keeps_run_number_new_file_becomes_run_two_with_automatic_comparison(self):
        kw = dict(use_mock=True, base_dir=self.ws, build_decks=False)
        r1 = run_pipeline(tb_file=self.day3, **kw)
        again = run_pipeline(tb_file=self.day3, **kw)
        r2 = run_pipeline(tb_file=self.day4, **kw)
        r2_again = run_pipeline(tb_file=self.day4, **kw)
        self.assertEqual((r1["run_no"], again["run_no"], r2["run_no"], r2_again["run_no"]), (1, 1, 2, 2))
        self.assertEqual(len(r2["_frames"]["revision_log"]), 6)
        self.assertEqual(len(r2_again["_frames"]["revision_log"]), 6)       # a re-run keeps the same comparison
        self.assertTrue((self.paths.out_dir("2026-06_r2") / "revision_log_INTERNAL.csv").exists())


# ======================================================================= stage 5-6: packs, exclusions, redaction
class TestPacksAndSafeguards(WorkspaceCase):
    def setUp(self):
        super().setUp()
        self.r = self.run_day4(write_outputs=True)
        self.texts = self.r["_frames"]["pack_texts"]

    def test_redaction_levels(self):
        red = Redactor(read_table(self.paths.config / "entity_master.csv", ["Entity_Type", "Legal_Name", "Aliases"]))
        text = "Orchid Bay Holdings Pte Ltd paid via 088-123456-7; contact Jane Tan jane.tan@x.example +65 8123 4567 S1234567D; unknown Foo Bar Pte Ltd"
        self.assertEqual(red.redact(text, "none")[0], text)
        pii, _ = red.redact(text, "pii")
        self.assertIn("Orchid Bay", pii)
        for secret in ("088-123456-7", "Jane Tan", "jane.tan@", "+65 8123", "S1234567D"):
            self.assertNotIn(secret, pii)
        full, counts = red.redact(text, "full")
        self.assertNotIn("Orchid Bay", full); self.assertNotIn("Foo Bar Pte Ltd", full)
        self.assertEqual(red.find_leaks(full, "full"), []); self.assertGreaterEqual(counts["CUSTOMER"], 1)

    def test_investor_material_and_deck_contain_no_identifier_management_keeps_them(self):
        investor = self.texts["Investor"] + self.r["narratives"]["Investor"] + json.dumps(self.r["slides"]["Investor"])
        for secret in SEEDED:
            self.assertNotIn(secret, investor, f"{secret} leaked")
        self.assertIn("Orchid Bay", self.texts["Management"])
        self.assertGreater(sum(self.r["protection_log"]["Investor"]["redaction_counts"].values()), 0)

    def test_board_is_pii_redacted_but_customer_names_remain(self):
        for secret in ("088-123456-7", "Jane Tan", "jane.tan@"):
            self.assertNotIn(secret, self.texts["Board"])
        self.assertIn("Orchid Bay Holdings", self.texts["Board"])

    def test_external_packs_do_not_reveal_management_only_concepts(self):
        for ext in ("Board", "Investor"):
            low = (self.texts[ext] + self.r["narratives"][ext]).lower()
            for word in ("adjusted gross margin", "apportioned", *[b.lower() for b in BU_NAMES]):
                self.assertNotIn(word, low, f"{word} in {ext}")
        self.assertIn("Adjusted gross margin", self.texts["Management"])

    def test_investor_policy_budget_qoq_yoy_kpis_totals_only(self):
        inv = self.texts["Investor"]
        for horizon in ("BvA_Month", "BvA_QTD", "BvA_YTD", "QoQ", "YoY"):
            self.assertIn(horizon, inv)
        for horizon in ("RFvA_", "MoM"):
            self.assertNotIn(horizon, inv)
        self.assertNotIn("rolling", inv.lower()); self.assertNotIn("forecast", inv.lower())
        for line in ("Product Revenue", "Service Revenue", "Cost of Revenue", "R&D Expense", "Sales & Marketing"):
            self.assertNotIn(f"| {line} |", inv)                         # totals only
        self.assertIn("| Total Revenue |", inv)
        self.assertIn("KPI scorecard", inv); self.assertNotIn("Cash balance", inv)
        self.assertNotIn("Comparisons - Cash flow", inv); self.assertNotIn("Comparisons - Balance sheet", inv)
        self.assertIn("Comparisons - Cash flow", self.texts["Board"]); self.assertIn("Comparisons - Balance sheet", self.texts["Management"])

    def test_exclusion_rules_are_data_not_code(self):
        rules = self.ws / "config" / "exclusion_rules.csv"
        rules.write_text("\n".join(l for l in rules.read_text().splitlines() if ",MoM," not in l) + "\n")
        self.assertIn("MoM", self.run_day4()["_frames"]["pack_texts"]["Investor"])

    def test_revision_log_never_enters_a_pack_board_and_management_get_one_line(self):
        for text in self.texts.values():
            self.assertNotIn("ADJ-0", text); self.assertNotIn("Prepared_By", text)
        self.assertIn("Figures were revised", self.texts["Board"]); self.assertIn("Figures were revised", self.texts["Management"])
        self.assertNotIn("Figures were revised", self.texts["Investor"])

    def test_raw_ledger_memos_stay_internal_and_notes_follow_external_flag(self):
        self.assertIn("[Ledger memo]", self.texts["Management"])
        for ext in ("Board", "Investor"):
            self.assertNotIn("[Ledger memo]", self.texts[ext]); self.assertNotIn("SMB launch", self.texts[ext])   # note N3 is External_OK = No

    def test_notes_todo_lists_material_lines_without_a_note(self):
        p = self.ws / "data" / "input" / "driver_notes.csv"
        p.write_text("\n".join(l for l in p.read_text().splitlines() if "Sales & Marketing" not in l) + "\n")
        r = self.run_day4()
        self.assertEqual(r["notes_todo"][0]["Line"], "Sales & Marketing")
        self.assertTrue(any("Add a driver note for Sales & Marketing" in i["Item"] for i in r["review_list"]))

    def test_output_files_and_manifest(self):
        out = self.paths.out_dir(self.r["version_id"])
        for name in ("packs/management_pack.md", "packs/board_pack.md", "packs/investor_pack.md", "ai_output.json", "revision_log_INTERNAL.csv",
                     "input_manifest.json", "signoff_investor.json", "bridge.csv", "kpi_scorecard.csv", "review_list.json", "trial_balance_snapshot.csv"):
            self.assertTrue((out / name).exists(), name)
        manifest = json.loads((out / "input_manifest.json").read_text())
        self.assertTrue({"analyst_skill.md", "audience_cards.md", "standard_prompt.md", "actuals_history.csv"} <= set(manifest))


class TestSkillAndPrompt(unittest.TestCase):
    def test_skill_and_cards_respect_character_limits(self):
        paths = Paths(ROOT)
        self.assertLessEqual(len(paths.skill.read_text()), SKILL_MAX_CHARS)
        for section in paths.cards.read_text().split("## Audience: ")[1:]:
            self.assertLessEqual(len(section), CARD_MAX_CHARS)

    def test_standard_prompt_specifies_the_json_schema(self):
        text = Paths(ROOT).prompt.read_text()
        for key in ("commentary", "analysis", "key_variances", "unexplained", "slides", "layout", "bullets"):
            self.assertIn(key, text)


# ======================================================================= stage 7-8: AI output and checks
class TestAiOutputAndChecks(WorkspaceCase):
    def test_ai_output_from_another_platform_goes_through_the_same_controls(self):
        base = self.run_day4(write_outputs=False)["_frames"]["outputs"]["Investor"]
        good = self.ws / "inv.json"
        good.write_text(json.dumps({k: base[k] for k in ("commentary", "analysis", "slides")}))
        r = self.run_day4(ai_override={"Investor": good})
        self.assertEqual((r["status"], r["narrative_modes"]["Investor"]), ("SUCCESS", "EXTERNAL"))
        leaky = {k: base[k] for k in ("commentary", "analysis", "slides")} | {"commentary": base["commentary"] + " Orchid Bay Holdings Pte Ltd said revenue was $999,999 versus budget."}
        good.write_text(json.dumps(leaky))
        r = self.run_day4(ai_override={"Investor": good})
        self.assertFalse(r["checks"]["Investor"]["passed"])                              # invented figure and identifier are caught
        self.assertNotIn("Orchid Bay", r["narratives"]["Investor"])                      # and redacted before anyone sees it
        good.write_text("not json")
        self.assertEqual(self.run_day4(ai_override={"Investor": good})["status"], "FAILED")

    def test_mock_output_is_data_driven_and_passes_all_checks(self):
        r = self.run_day4()
        for audience, check in r["checks"].items():
            self.assertTrue(check["passed"], f"{audience}: {check['numbers']['unmatched']} {check['forbidden_terms']} {check['spec_problems']}")
        self.assertIn("-$141,800", r["narratives"]["Management"])
        self.assertNotIn("software expansion", r["narratives"]["Investor"])

    def test_number_check_catches_invented_figures(self):
        pack = "Operating Profit 28,500 variance -141,500 (83.2%) DSO 44.80 days 2026"
        self.assertTrue(check_numbers("Operating Profit was $28,500, down 83.2%.", pack)["passed"])
        bad = check_numbers("Operating Profit was $28,500 and revenue grew $999,999 (12.3%).", pack)
        self.assertEqual(sorted(bad["unmatched"]), ["$999,999", "12.3%"])

    def test_forbidden_terms(self):
        self.assertEqual(check_forbidden_terms("Results beat the Budget and our forecast.", ["budget", "forecast", "guidance"]), ["budget", "forecast"])

    def test_slide_spec_limits(self):
        good = {"commentary": "x", "analysis": {k: [] for k in ("key_variances", "drivers", "unexplained", "questions")} | {"headline": "h"},
                "slides": [{"layout": "headline", "title": "T", "bullets": ["a"]}]}
        self.assertEqual(check_slide_spec(good, "Investor"), [])
        bad = {**good, "slides": [{"layout": "cash", "title": "T" * 80, "bullets": ["a"] * 5}] * 8}
        problems = " | ".join(check_slide_spec(bad, "Investor"))
        for expected in ("exceeds the limit", "not allowed for Investor", "longer than 70", "more than 4 bullets"):
            self.assertIn(expected, problems)

    def test_live_failure_is_labelled_fallback_not_silent(self):
        """Original bug: Ollama errors silently returned mock text that looked like AI output."""
        r = self.run_day4(use_mock=False, ollama_url="http://127.0.0.1:9")
        self.assertEqual(r["status"], "SUCCESS"); self.assertEqual(set(r["narrative_modes"].values()), {"MOCK_FALLBACK"})
        self.assertTrue(r["narratives"]["Board"].startswith("[FALLBACK")); self.assertEqual(r["stages"][6]["status"], "WARN")

    def test_live_request_sets_json_mode_and_context_window(self):
        sent = {}

        def fake_post(url, json=None, timeout=None):
            sent.update(json)
            resp = mock.Mock(); resp.raise_for_status = lambda: None
            resp.json = lambda: {"response": '```json\n{"commentary": "ok", "analysis": {}, "slides": []}\n```'}
            return resp
        with mock.patch.object(requests, "post", fake_post):
            out = llm_narrative.generate_output({"comparisons": pd.DataFrame()}, "prompt", False, num_ctx=8192)
        self.assertEqual((out["mode"], sent["format"], sent["options"]["num_ctx"]), ("LIVE", "json", 8192))

    def test_invalid_json_is_retried_once_then_falls_back(self):
        calls = []
        resp = mock.Mock(); resp.raise_for_status = lambda: None; resp.json = lambda: {"response": "not json"}
        with mock.patch.object(requests, "post", lambda *a, **k: calls.append(1) or resp):
            out = llm_narrative.generate_output(self.run_day4()["_frames"]["packs"]["Board"], "prompt", False)
        self.assertEqual((len(calls), out["mode"]), (2, "MOCK_FALLBACK"))

    def test_prompt_near_context_window_warns_with_a_concrete_floor_that_reserves_output_headroom(self):
        """num_ctx is a budget shared by the prompt AND the model's reply. The warning must say so and suggest a
        concrete minimum with headroom reserved for the output, not just 'raise it above the prompt size' - that
        was the earlier wording, and it invites raising num_ctx to barely more than the prompt, leaving the model
        no room to actually write its response."""
        r = self.run_day4(num_ctx=500)
        matches = [w for w in r["warnings"] if "context window" in w]
        self.assertTrue(matches)
        for w in matches:
            self.assertIn("prompt AND the reply together", w)
            prompt_tokens = int(w.split(" tokens")[0].split()[-1])
            suggested = int(w.split("at least ")[1].split(" ")[0])
            self.assertGreaterEqual(suggested, prompt_tokens + 2048)   # real headroom, not just "above the prompt"

    def test_parse_ai_json_tolerates_prose_around_the_object(self):
        """Ollama's format=json guarantees valid JSON, not that the WHOLE reply is just the object."""
        text = 'Sure, here is the analysis:\n{"commentary": "ok", "analysis": {"headline": "h"}, "slides": []}\nHope that helps!'
        data = parse_ai_json(text)
        self.assertEqual(data["commentary"], "ok")
        self.assertEqual(data["analysis"]["questions"], [])             # missing optional sub-key filled with []

    def test_parse_ai_json_fills_missing_optional_analysis_keys_but_still_requires_commentary_and_slides(self):
        data = parse_ai_json('{"commentary": "ok", "slides": [{"layout": "headline", "title": "T", "bullets": []}]}')
        self.assertEqual(data["analysis"], {"headline": "", "key_variances": [], "drivers": [], "unexplained": [], "questions": []})
        with self.assertRaises(Exception):
            parse_ai_json('{"analysis": {}, "slides": []}')             # commentary missing: still an error
        with self.assertRaises(Exception):
            parse_ai_json('{"commentary": "ok"}')                       # slides missing: still an error

    def test_extract_json_object_is_quote_aware(self):
        text = 'noise {"a": "a brace } inside a string", "b": 1} trailing'
        self.assertEqual(extract_json_object(text), '{"a": "a brace } inside a string", "b": 1}')

    def test_check_ollama_connection_and_model_match(self):
        import requests as req
        from unittest import mock as _mock
        ok = _mock.Mock(); ok.raise_for_status = lambda: None
        ok.json = lambda: {"models": [{"name": "llama3.2:latest"}, {"name": "qwen2.5:7b"}]}
        with _mock.patch.object(req, "get", lambda *a, **k: ok):
            probe = check_ollama_connection("http://localhost:11434")
        self.assertTrue(probe["ok"])
        self.assertTrue(model_is_available("llama3.2", probe["models"]))     # matches by base name, ignoring :tag
        self.assertFalse(model_is_available("mistral", probe["models"]))
        with _mock.patch.object(req, "get", side_effect=req.exceptions.ConnectionError("refused")):
            down = check_ollama_connection("http://localhost:11434")
        self.assertFalse(down["ok"])
        self.assertIn("ConnectionError", down["error"])

    def test_timeout_is_retried_but_connection_refused_is_not(self):
        import requests as req
        from unittest import mock as _mock
        calls = []

        def timeout_then_ok(*a, **k):
            calls.append(1)
            raise req.exceptions.Timeout("slow")
        with _mock.patch.object(req, "post", timeout_then_ok):
            out = llm_narrative.generate_output(self.run_day4()["_frames"]["packs"]["Board"], "prompt", False, retries=3, timeout=1)
        self.assertEqual(len(calls), 3)                                 # all 3 attempts used for a timeout
        self.assertIn("Timeout", out["error"])

        calls.clear()
        with _mock.patch.object(req, "post", side_effect=req.exceptions.ConnectionError("refused")):
            out2 = llm_narrative.generate_output(self.run_day4()["_frames"]["packs"]["Board"], "prompt", False, retries=3, timeout=1)
        self.assertEqual(out2["mode"], "MOCK_FALLBACK")
        self.assertIn("ConnectionError", out2["error"])


# ======================================================================= stage 9: deck
class TestDeck(WorkspaceCase):
    def test_deck_never_exceeds_max_slides_even_if_the_ai_returns_more(self):
        """A live model does not always respect 'at most 7 slides'. check_slide_spec flags that for review, but
        the deck ITSELF must never physically exceed the limit regardless of what the AI returned."""
        from pptx import Presentation
        from engine.config import MAX_SLIDES
        r = self.run_day4(write_outputs=False, build_decks=False)
        pack, out = r["_frames"]["packs"]["Board"], dict(r["_frames"]["outputs"]["Board"])
        extra = {**out["slides"][0], "title": "Extra slide"}
        out["slides"] = out["slides"] + [extra] * 5                     # force well over the limit
        self.assertGreater(len(out["slides"]), MAX_SLIDES)
        path = self.ws / "overlong.pptx"
        build_deck(pack, out, path)
        self.assertEqual(len(Presentation(path).slides), MAX_SLIDES + 1)  # +1 for the title slide
        problems = check_slide_spec(out, "Board")
        self.assertTrue(any("exceeds the limit" in p for p in problems))  # still flagged for FP&A review

    def test_decks_built_checked_and_contain_only_audience_content(self):
        from pptx import Presentation
        r = self.run_day4(write_outputs=True, build_decks=True)
        self.assertEqual(set(r["decks"]), {"Management", "Board", "Investor"})
        self.assertTrue(all(c["passed"] for c in r["deck_checks"].values()), r["deck_checks"])
        text = {a: " ".join(sh.text_frame.text for s in Presentation(p).slides for sh in s.shapes if sh.has_text_frame) for a, p in r["decks"].items()}
        self.assertIn("DRAFT run 2", text["Investor"])
        for secret in SEEDED + ["Operating cash flow", "Cash balance", "adjusted gross margin"]:
            self.assertNotIn(secret.lower(), text["Investor"].lower())
        self.assertEqual(len(Presentation(r["decks"]["Investor"]).slides), 3)

    def test_deck_check_detects_overflow_and_invented_numbers(self):
        r = self.run_day4()
        pack, out = r["_frames"]["packs"]["Investor"], json.loads(json.dumps(r["_frames"]["outputs"]["Investor"]))
        out["slides"][0]["bullets"] = ["x " * 1500, "Revenue was $999,999."]
        path = self.ws / "bad.pptx"
        build_deck(pack, out, path)
        res = check_deck(path, r["_frames"]["pack_texts"]["Investor"])
        self.assertTrue(res["overflow"]); self.assertIn("$999,999", res["numbers"]["unmatched"]); self.assertFalse(res["passed"])


# ======================================================================= stage 10-11: sign-off
class TestSignoff(unittest.TestCase):
    def test_management_needs_fd_and_cfo_in_either_order(self):
        s = new_state("v", "Management")
        self.assertEqual(sorted(pending_roles(s)), ["CFO", "Finance Director"])
        with self.assertRaises(SignoffError):
            approve(s, "CEO")
        approve(s, "CFO"); self.assertEqual(pending_roles(s), ["Finance Director"])
        approve(s, "Finance Director"); self.assertEqual(s.status, RELEASED)

    def test_investor_deck_also_needs_the_ceo_after_fd_and_cfo(self):
        s = new_state("v", "Investor")
        with self.assertRaises(SignoffError):
            approve(s, "CEO")
        approve(s, "Finance Director"); approve(s, "CFO")
        self.assertEqual((s.status, pending_roles(s)), ("PENDING", ["CEO"]))
        approve(s, "CEO"); self.assertEqual(s.status, RELEASED)

    def test_wording_rejection_routes_to_fpa_and_restarts(self):
        s = new_state("v", "Board"); approve(s, "Finance Director")
        reject(s, "CFO", "WORDING", "Tone too strong")
        self.assertEqual((s.status, s.approvals), (FPA_EDIT, [])); resubmit(s); self.assertEqual(s.status, "PENDING")

    def test_numbers_rejection_needs_a_new_run(self):
        s = new_state("2026-06_r1", "Management")
        reject(s, "Finance Director", "NUMBERS", "Revenue not fully closed")
        self.assertEqual(s.status, FINANCE_REVIEW)
        with self.assertRaises(SignoffError):
            resubmit(s)
        resubmit(s, "2026-06_r2"); self.assertEqual((s.version_id, s.status), ("2026-06_r2", "PENDING"))

    def test_rejection_requires_type_reason_and_is_logged(self):
        s = new_state("v", "Management")
        for args in (("Finance Director", "WORDING", "  "), ("Finance Director", "TYPO", "x"), ("CEO", "WORDING", "x")):
            with self.assertRaises(SignoffError):
                reject(s, *args)
        reject(s, "Finance Director", "NUMBERS", "late accrual")
        self.assertEqual((s.log[-1]["Reason"], s.log[-1]["Reject_Type"]), ("late accrual", "NUMBERS"))

    def test_edit_after_approval_resets_signoff(self):
        s = new_state("v", "Management"); approve(s, "Finance Director"); approve(s, "CFO")
        record_edit(s, "FP&A", "Changed a sentence"); self.assertEqual((s.status, s.approvals), ("PENDING", []))


# ======================================================================= roll-forward, history, Sheets
def july_export(day4_path: Path, out_path: Path) -> None:
    """Build a realistic July ERP export from June: P&L x1.10, balance sheet rolled forward, cash as the plug."""
    df = pd.read_excel(day4_path, dtype=str)
    df["Balance"] = df["Balance"].astype(float)
    code = lambda c: df.loc[df["Account Code"] == c, "Balance"]  # noqa: E731
    pl = df["Account Code"].str[0].isin(["4", "5", "6"])
    june_ni = code("4000-01").sum() + code("4000-02").sum() - code("5000-10").sum() - code("6000-05").sum() - code("6000-20").sum() - code("6000-90").sum() - code("6000-80").sum()
    df.loc[pl, "Balance"] = (df.loc[pl, "Balance"] * 1.10).round(-2)
    july = lambda c: df.loc[df["Account Code"] == c, "Balance"].sum()  # noqa: E731
    ni = july("4000-01") + july("4000-02") - sum(july(c) for c in ("5000-10", "6000-05", "6000-20", "6000-90", "6000-80"))
    new = {"3100-00": code("3100-00").iloc[0] + june_ni, "1500-00": code("1500-00").iloc[0] + 20000 - july("6000-80")}
    for c, v in new.items():
        df.loc[df["Account Code"] == c, "Balance"] = v
    cash = july("2100-00") + july("2200-00") + july("3000-00") + july("3100-00") + ni - july("1100-00") - july("1200-00") - july("1500-00")
    df.loc[df["Account Code"] == "1000-00", "Balance"] = cash
    df["Journal Description"] = None
    df.to_excel(out_path, index=False)


class TestRollForward(WorkspaceCase):
    def release_june(self):
        kw = dict(use_mock=True, base_dir=self.ws, build_decks=False)
        run_pipeline(tb_file=self.day3, **kw)
        r2 = run_pipeline(tb_file=self.day4, **kw)
        return r2["version_id"]

    def test_one_click_approvals_release_and_lock_automatically(self):
        vid = self.release_june()
        self.assertEqual(vid, "2026-06_r2")
        first = signoff_action(vid, "Finance Director", base_dir=self.ws)
        self.assertEqual(set(first["acted_on"]), {"Management", "Board", "Investor"}); self.assertNotIn("lock", first)
        signoff_action(vid, "CFO", base_dir=self.ws)
        with self.assertRaises(SignoffError):
            lock_after_release(vid, self.ws)                      # investor deck still waits for the CEO
        final = signoff_action(vid, "CEO", base_dir=self.ws)
        self.assertEqual(final["acted_on"], ["Investor"]); self.assertEqual(final["lock"]["rows_appended"], 29)
        self.assertIn("2026-06", set(pd.read_csv(self.paths.history, dtype=str)["Period"]))

    def test_rerunning_a_locked_month_is_refused(self):
        approve_all(self.release_june(), self.ws)
        r = self.run_day4()
        self.assertEqual(r["status"], "FAILED"); self.assertIn("locked", r["error"])

    def test_next_month_starts_from_fresh_file_and_compares_to_locked_june(self):
        approve_all(self.release_june(), self.ws)
        july = self.ws / "july.xlsx"
        july_export(self.day4, july)
        r = run_pipeline(tb_file=july, period="2026-07", use_mock=True, base_dir=self.ws, write_outputs=False, build_decks=False)
        self.assertEqual(r["status"], "SUCCESS", r["error"])
        cmp_df = r["_frames"]["stat_cmp"]
        mom = cmp_df[(cmp_df["Statement"] == "P&L") & (cmp_df["Horizon"] == "MoM") & (cmp_df["Line"] == "Total Revenue")].iloc[0]
        self.assertAlmostEqual(mom["Comparator"], 608000.0); self.assertAlmostEqual(mom["Var_Amt"], 60800.0)

    def test_restatement_needs_reason_and_finance_director_and_is_logged(self):
        tb = load_trial_balance(self.day4, "2026-06", load_column_map(self.paths.config / "erp_column_map.csv"), [])
        lock_period(self.paths.history, tb, "2026-06", "r2")
        fixed = tb.copy(); fixed.loc[0, "Amount"] += 1000
        log, hist = self.paths.restatement_log, self.paths.history
        for args in (("", "Finance Director"), ("Error found", "CFO")):
            with self.assertRaises(HistoryError):
                restate_period(hist, log, fixed, "2026-06", args[0], args[1], "A")
        with self.assertRaises(HistoryError):
            lock_period(hist, tb, "2026-06", "r2")
        changed = restate_period(hist, log, fixed, "2026-06", "Error found after close", "Finance Director", "A")
        self.assertEqual(len(changed), 1); self.assertAlmostEqual(changed.iloc[0]["Change"], 1000); self.assertEqual(len(pd.read_csv(log)), 1)


class TestGoogleSheetsCompatibility(WorkspaceCase):
    def test_sheets_workbook_gives_identical_results_to_csv_inputs(self):
        csv_run = self.run_day4()
        sheets_run = self.run_day4(sheets_workbook=self.paths.input / "sheets_inputs.xlsx")
        self.assertEqual(sheets_run["status"], "SUCCESS", sheets_run["error"])
        for key in ("ratios", "material_by_horizon", "narratives", "kpis"):
            self.assertEqual(csv_run[key], sheets_run[key], key)

    def test_sheets_date_conversion_and_stray_spaces_are_tolerated(self):
        path = self.paths.input / "sheets_inputs.xlsx"
        wb = load_workbook(path); notes = wb["Driver_Notes"]
        for row in range(2, notes.max_row + 1):
            notes.cell(row, 2).value = "2026-06-01 00:00:00"
            notes.cell(row, 5).value = " " + str(notes.cell(row, 5).value) + " "
        wb.save(path)
        csv_run, r = self.run_day4(), self.run_day4(sheets_workbook=path)
        self.assertEqual(r["status"], "SUCCESS", r["error"])
        self.assertTrue(any("looked like a date" in w for w in r["warnings"]))
        self.assertEqual(r["narratives"]["Management"], csv_run["narratives"]["Management"])

    def test_kpi_targets_edited_in_sheets_change_the_scorecard(self):
        path = self.paths.input / "sheets_inputs.xlsx"
        wb = load_workbook(path); ws = wb["KPI_Targets"]
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, 1).value == "Gross_Margin_%":
                ws.cell(row, 3).value = "60"
        wb.save(path)
        status = lambda r: [k for k in r["kpis"] if k["KPI"] == "Gross_Margin_%"][0]["Status"]  # noqa: E731
        self.assertEqual(status(self.run_day4()), "Watch"); self.assertEqual(status(self.run_day4(sheets_workbook=path)), "On target")

    def test_missing_tab_gives_clear_error(self):
        path = self.paths.input / "sheets_inputs.xlsx"
        wb = load_workbook(path); del wb["Driver_Notes"]; wb.save(path)
        r = self.run_day4(sheets_workbook=path)
        self.assertEqual(r["status"], "FAILED"); self.assertIn("Driver_Notes", r["error"])


class TestFormatting(unittest.TestCase):
    def test_formats(self):
        self.assertEqual(fmt.money(-141500), "-$141,500"); self.assertEqual(fmt.money(float("nan")), "n/a")
        self.assertEqual(fmt.pct_fraction(0.0789), "7.9%"); self.assertEqual(fmt.days(44.8), "44.80 days")
        self.assertEqual(fmt.period_label("2026-06"), "June 2026")


if __name__ == "__main__":
    unittest.main()
