"""Generate deterministic MOCK inputs for the showcase. All names, emails and account numbers are fake.

Run from the project root:   python data/generate_mock_data.py

A small three-statement model keeps the mock internally consistent: the P&L drives the balance sheet
(receivables, payables, fixed assets, equity) and cash is the balancing item, so the cash flow the
engine derives reconciles exactly to the change in cash.

Creates
  data/history/actuals_history.csv      18 locked months (Dec 2024 - May 2026): P&L and balance sheet
  data/input/trial_balance.xlsx         June 2026 Day-3 ERP export (the file Finance drops each run)
  data/input/samples/day3_trial_balance.xlsx / day4_trial_balance.xlsx   the two sample exports
  data/input/budget.xlsx, forecast.xlsx P&L and balance sheet by month
  data/input/driver_notes.csv, revision_reasons.csv, kpi_targets.csv
  data/input/sheets_inputs.xlsx         Google Sheets-ready workbook of the human-owned tables
"""
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from engine.variance_ratios import days_in  # noqa: E402

SEED, REPORT = 42, "2026-06"
BUS = ["Platform", "Enterprise", "SMB"]
PL = ["4000-01", "4000-02", "5000-10", "6000-05", "6000-20", "6000-90", "6000-80"]
BS = ["1000-00", "1100-00", "1200-00", "1500-00", "2100-00", "2200-00", "3000-00", "3100-00"]
DESC = {"4000-01": "Gross Software Revenue", "4000-02": "Consulting Services Income", "5000-10": "Direct Hosting & Cloud Expense",
        "6000-05": "Salaries - Engineering & Product", "6000-20": "Digital Marketing & Ad Spend",
        "6000-90": "Executive & Admin Overhead", "6000-80": "Depreciation", "1000-00": "Cash and Bank",
        "1100-00": "Trade Receivables", "1200-00": "Prepaid Expenses", "1500-00": "Fixed Assets (net)",
        "2100-00": "Trade Payables", "2200-00": "Accrued Expenses", "3000-00": "Share Capital",
        "3100-00": "Retained Earnings (opening)"}
BS_KEY = {"1000-00": "cash", "1100-00": "ar", "1200-00": "prepaid", "1500-00": "fa", "2100-00": "ap",
          "2200-00": "accr", "3000-00": "sc", "3100-00": "re_open"}
SHARES = {"4000-01": [.50, .35, .15], "4000-02": [.10, .60, .30], "5000-10": [.60, .30, .10], "6000-05": [.55, .30, .15],
          "6000-20": [.30, .30, .40], "6000-90": [.34, .33, .33], "6000-80": [.34, .33, .33]}
BUD_SHARES = {**SHARES, "4000-01": [.48, .37, .15], "4000-02": [.12, .58, .30], "5000-10": [.58, .32, .10], "6000-20": [.34, .33, .33]}
JUNE_V1 = {"4000-01": 520000, "4000-02": 110000, "5000-10": 185000, "6000-05": 160000, "6000-20": 145000, "6000-90": 75000, "6000-80": 8300}
ANCHOR_MAY26 = {"4000-01": 505000, "4000-02": 128000, "5000-10": 172000, "6000-05": 155000, "6000-20": 118000, "6000-90": 72000, "6000-80": 8100}
BUDGET_JUNE = {"4000-01": 500000, "4000-02": 150000, "5000-10": 150000, "6000-05": 150000, "6000-20": 110000, "6000-90": 70000, "6000-80": 8000}
FORECAST_JUNE = {"4000-01": 510000, "4000-02": 120000, "5000-10": 170000, "6000-05": 155000, "6000-20": 130000, "6000-90": 72000, "6000-80": 8200}
GROWTH = {"4000-01": .016, "4000-02": .004, "5000-10": .012, "6000-05": .008, "6000-20": .010, "6000-90": .004, "6000-80": .004}
OPENING = {"sc": 4_000_000, "re_open": -2_500_000, "fa": 300_000}     # balance sheet at Dec 2024
RAISE = {"2026-02": 500_000}                                          # equity raise (financing cash flow)


def split(total: float, shares: list) -> list:
    """Split a total across BUs, rounded to 100, remainder on the largest BU so it ties exactly."""
    parts = [round(total * s, -2) for s in shares]
    parts[int(np.argmax(shares))] += round(total, -2) - sum(parts)
    return parts


def months(start: str, end: str) -> list:
    """Inclusive list of YYYY-MM periods."""
    return list(pd.period_range(start, end, freq="M").strftime("%Y-%m"))


def k_from(period: str, anchor: str) -> int:
    """Whole months between a period and an anchor month."""
    return (int(period[:4]) - int(anchor[:4])) * 12 + int(period[5:7]) - int(anchor[5:7])


def run_model(pl: dict, period_list: list, opening: dict, dso, dpo, capex: dict, overrides: dict = None) -> dict:
    """Roll a balance sheet forward from the P&L; cash is the balancing item. Returns {period: state}."""
    overrides, out, prev = overrides or {}, {}, None
    for m in period_list:
        p = pl[m]
        rev, cogs, dep = p["4000-01"] + p["4000-02"], p["5000-10"], p["6000-80"]
        cash_opex = p["6000-05"] + p["6000-20"] + p["6000-90"]
        ni = rev - cogs - cash_opex - dep
        s = dict(opening) if prev is None else {"sc": prev["sc"] + RAISE.get(m, 0), "re_open": prev["re_open"] + prev["ni"],
                                                "fa": prev["fa"] + capex[m] - dep}
        s["ar"], s["ap"] = round(dso(m) * rev / days_in([m]), -2), round(dpo(m) * cogs / days_in([m]), -2)
        s["prepaid"], s["accr"] = round(.08 * cash_opex, -2), round(.12 * cash_opex, -2)
        s.update(overrides.get(m, {}))
        s["ni"] = ni
        s["cash"] = s["ap"] + s["accr"] + s["sc"] + s["re_open"] + ni - s["ar"] - s["prepaid"] - s["fa"]
        out[m] = prev = s
    return out


def rows_for(pl: dict, states: dict, shares: dict) -> list:
    """Long-format rows (Period, code, BU, Amount): P&L split by BU plus balance sheet on 'Corporate'."""
    rows = []
    for m in states:
        for a in PL:
            rows += [(m, a, bu, amt) for bu, amt in zip(BUS, split(pl[m][a], shares[a]))]
        rows += [(m, code, "Corporate", states[m][key]) for code, key in BS_KEY.items()]
    return rows


def build_models(rng):
    """Actual (through June v1), budget and forecast P&L + balance sheets."""
    act_months = months("2024-12", "2026-06")
    act_pl = {}
    for m in act_months:
        if m == REPORT:
            act_pl[m] = dict(JUNE_V1)
            continue
        act_pl[m] = {a: round(ANCHOR_MAY26[a] * (1 + GROWTH[a]) ** k_from(m, "2026-05") * (1 + (0 if a == "6000-80" else rng.normal(0, .015))), -2)
                     for a in PL}
    capex_a = {m: round(rng.uniform(15000, 35000), -2) for m in act_months}
    dso_a = {m: 40 + 3.5 * i / 18 + rng.normal(0, .8) for i, m in enumerate(act_months)}
    dpo_a = {m: 30 - 2 * i / 18 + rng.normal(0, .6) for i, m in enumerate(act_months)}
    actual = run_model(act_pl, act_months, OPENING, dso_a.get, dpo_a.get, capex_a, {REPORT: {"ar": 930000, "ap": 172000}})
    bud_months = months("2024-12", "2026-12")
    bud_pl = {m: {a: round(BUDGET_JUNE[a] * (1 + GROWTH[a]) ** k_from(m, REPORT), -2) for a in PL} for m in bud_months}
    budget = run_model(bud_pl, bud_months, OPENING, lambda m: 40, lambda m: 30, {m: 25000 for m in bud_months})
    f_months = months("2025-12", "2027-05")
    f_pl = {}
    for m in f_months:
        if m < REPORT:
            f_pl[m] = {a: round(act_pl[m][a] * (1 + (0 if m == "2025-12" or a == "6000-80" else rng.normal(0, .02))), -2) for a in PL}
        else:
            f_pl[m] = {a: round(FORECAST_JUNE[a] * (1 + GROWTH[a]) ** k_from(m, REPORT), -2) for a in PL}
    opening_f = {k: actual["2025-12"][k] for k in ("sc", "re_open", "fa")}
    forecast = run_model(f_pl, f_months, opening_f, lambda m: 43, lambda m: 29, {m: 25000 for m in f_months})
    return act_pl, actual, bud_pl, budget, f_pl, forecast


def tb_frame(pl_m: dict, state: dict, memos: dict) -> pd.DataFrame:
    """One month as an ERP-style trial balance export (ERP headers, no Period/Version columns)."""
    rows = []
    for a in PL:
        rows += [(a, DESC[a], bu, amt, memos.get((a, bu), "")) for bu, amt in zip(BUS, split(pl_m[a], SHARES[a]))]
    rows += [(code, DESC[code], "Corporate", state[key], memos.get((code, "Corporate"), "")) for code, key in BS_KEY.items()]
    return pd.DataFrame(rows, columns=["Account Code", "Account Name", "Business Unit", "Balance", "Journal Description"])


MEMOS_V1 = {("4000-01", "Enterprise"): "Renewal invoicing Sunrise Logistics Pte Ltd; remittance to account 088-123456-7",
            ("4000-02", "Enterprise"): "Phase 2 consulting for Orchid Bay Holdings Pte Ltd; contact Jane Tan +65 8123 4567",
            ("5000-10", "Platform"): "Hosting invoices from CloudNimbus Hosting Pte Ltd during region migration"}
CHANGES_V2 = [("4000-01", "Enterprise", -22000, "Cut-off correction: Sunrise Logistics Pte Ltd renewal go-live slipped into July"),
              ("1100-00", "Corporate", -22000, "Receivable reduced for cut-off correction"),
              ("5000-10", "Platform", +14500, "Accrual for CloudNimbus Hosting Pte Ltd June invoice INV-8841; billing contact ap@cloudnimbus.example"),
              ("2200-00", "Corporate", +14500, "Accrued hosting liability for June invoice"),
              ("6000-90", "Enterprise", -12000, "Meridian Health Corp event sponsorship coded to G&A in error"),
              ("6000-20", "Enterprise", +12000, "Recoded to marketing: Meridian Health Corp event sponsorship")]


def day4_from_day3(day3: pd.DataFrame) -> pd.DataFrame:
    """Day-4 export: Day-3 plus the three finance adjustments (revenue cut-off, missed accrual, miscoding)."""
    tb = day3.copy()
    for code, bu, delta, memo in CHANGES_V2:
        mask = (tb["Account Code"] == code) & (tb["Business Unit"] == bu)
        tb.loc[mask, "Balance"] += delta
        tb.loc[mask, "Journal Description"] = memo
    return tb


DRIVER_NOTES = pd.DataFrame([
    ("N1", REPORT, "Service Revenue", "Shortfall because Orchid Bay Holdings Pte Ltd deferred its phase 2 go-live to the next quarter; account contact Jane Tan (jane.tan@orchidbay.example).", "Yes"),
    ("N2", REPORT, "Cost of Revenue", "Higher cost reflects dual-running charges from CloudNimbus Hosting Pte Ltd during the region migration.", "Yes"),
    ("N3", REPORT, "Sales & Marketing", "Campaign spend was pulled forward from the next quarter for the SMB launch.", "No"),
    ("N4", REPORT, "Management basis", "Adjusted gross margin apportions part of engineering payroll and admin overhead to cost of revenue; for management use only.", "No"),
    ("N5", REPORT, "Product Revenue", "A large renewal with Sunrise Logistics Pte Ltd was signed late; payment to bank account 088-123456-7 is expected next week.", "Yes"),
], columns=["Note_ID", "Period", "Applies_To", "Note", "External_OK"])
# Optional manual overrides; most changed lines are explained automatically from the ERP journal description.
REVISION_REASONS = pd.DataFrame([(REPORT, "4000-01", "Enterprise", "ADJ-01", "Revenue cut-off: go-live slipped into July, revenue reversed", "Finance - Revenue Accounting")],
                                columns=["Period", "ERP_Account_Code", "BU", "Adjustment_Ref", "Reason", "Prepared_By"])
KPI_TARGETS = pd.DataFrame([
    ("Gross_Margin_%", "Higher", 70, 5, "Management|Board|Investor"),
    ("Adjusted_Gross_Margin_%", "Higher", 60, 5, "Management"),
    ("Operating_Margin_%", "Higher", 5, 40, "Management|Board|Investor"),
    ("Opex_%_of_Revenue", "Lower", 65, 5, "Management|Board|Investor"),
    ("Adjusted_Opex_%_of_Revenue", "Lower", 55, 5, "Management"),
    ("Revenue_Growth_QoQ_%", "Higher", 1.5, 30, "Management|Board|Investor"),
    ("Revenue_Growth_YoY_%", "Higher", 10, 20, "Management|Board|Investor"), ("DSO_Days", "Lower", 45, 5, "Management|Board|Investor"),
    ("DPO_Days", "Higher", 28, 10, "Management|Board|Investor"), ("Working_Capital_Drag_Days", "Lower", 18, 10, "Management|Board|Investor"),
    ("Cash_Balance", "Higher", 1000000, 10, "Management|Board"), ("Operating_Cash_Flow", "Higher", 50000, 20, "Management|Board"),
    ("Free_Cash_Flow", "Higher", 25000, 20, "Management|Board"), ("Runway_Months", "Higher", 18, 10, "Management|Board"),
], columns=["KPI", "Direction", "Target", "Tolerance_Pct", "Visible_To"])


def style_workbook(path: Path) -> None:
    """Light formatting for people: Arial font, bold header, frozen header row, readable widths."""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(min(max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2, 70), 10)
    wb.save(path)


def write_sheets_workbook(path: Path) -> None:
    """Workbook of the human-owned tables for Google Sheets (File > Import). Every cell is text so Sheets cannot convert codes or periods."""
    cfg = ROOT / "config"
    readme = pd.DataFrame({"README - Google Sheets input workbook": [
        "MOCK DATA ONLY: use synthetic data in a personal Google account; real data belongs in a company tenant.",
        "Tabs: Mapping_Statutory (Finance owns), Mapping_Management and Reclass_Rules (FP&A owns), Driver_Notes, Revision_Reasons, KPI_Targets.",
        "Keep Period as text (YYYY-MM) and ERP_Account_Code as text. In Sheets: Format > Number > Plain text before pasting.",
        "Do not rename tabs or columns. Do not add formulas; paste values only.",
        "To use: File > Download > Microsoft Excel (.xlsx), then run: python main.py --sheets-workbook <file>.xlsx",
        "Driver_Notes: External_OK = Yes lets the note reach the Board and investors (after redaction); No keeps it internal.",
        "Locked history, exclusion rules, entity master and sign-off state are NOT in Sheets by design."]})
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        for tab, frame in (("Mapping_Statutory", pd.read_csv(cfg / "coa_statutory.csv", dtype=str)),
                           ("Mapping_Management", pd.read_csv(cfg / "coa_management.csv", dtype=str)),
                           ("Reclass_Rules", pd.read_csv(cfg / "mgmt_reclass_rules.csv", dtype=str)),
                           ("Driver_Notes", DRIVER_NOTES), ("Revision_Reasons", REVISION_REASONS),
                           ("KPI_Targets", KPI_TARGETS.astype(str))):
            frame.to_excel(writer, sheet_name=tab, index=False)
    style_workbook(path)


def main() -> None:
    """Write every mock input file."""
    rng = np.random.default_rng(SEED)
    inp, hist_dir = ROOT / "data" / "input", ROOT / "data" / "history"
    (inp / "samples").mkdir(parents=True, exist_ok=True)
    hist_dir.mkdir(parents=True, exist_ok=True)
    act_pl, actual, bud_pl, budget, f_pl, forecast = build_models(rng)
    hist_months = [m for m in actual if m < REPORT]
    history = pd.DataFrame(rows_for(act_pl, {m: actual[m] for m in hist_months}, SHARES), columns=["Period", "ERP_Account_Code", "BU", "Amount"])
    history["Locked_Version"], history["Locked_On"] = "locked", "2026-06-05T10:00:00"
    history.to_csv(hist_dir / "actuals_history.csv", index=False)
    pd.DataFrame(rows_for(bud_pl, budget, BUD_SHARES), columns=["Period", "ERP_Account_Code", "BU", "Amount"]).to_excel(inp / "budget.xlsx", index=False)
    fc = pd.DataFrame(rows_for(f_pl, forecast, SHARES), columns=["Period", "ERP_Account_Code", "BU", "Amount"])
    fc["Forecast_Version"] = "2026-05"
    fc.to_excel(inp / "forecast.xlsx", index=False)
    day3 = tb_frame(act_pl[REPORT], actual[REPORT], MEMOS_V1)
    day3.to_excel(inp / "samples" / "day3_trial_balance.xlsx", index=False)
    day4_from_day3(day3).to_excel(inp / "samples" / "day4_trial_balance.xlsx", index=False)
    shutil.copy(inp / "samples" / "day3_trial_balance.xlsx", inp / "trial_balance.xlsx")
    DRIVER_NOTES.to_csv(inp / "driver_notes.csv", index=False)
    REVISION_REASONS.to_csv(inp / "revision_reasons.csv", index=False)
    KPI_TARGETS.to_csv(inp / "kpi_targets.csv", index=False)
    for f in ("budget", "forecast", "trial_balance", "samples/day3_trial_balance", "samples/day4_trial_balance"):
        style_workbook(inp / f"{f}.xlsx")
    write_sheets_workbook(inp / "sheets_inputs.xlsx")
    print("Mock data generated: 18 months history, Day-3/Day-4 ERP exports, budget, forecast, notes, KPI targets, Sheets workbook.")


if __name__ == "__main__":
    main()
